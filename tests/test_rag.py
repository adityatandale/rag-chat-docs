"""Offline tests: HashingEmbedder + in-memory Chroma + fake LLM. No network, no API key."""
import glob
import os

import numpy as np
import pytest

from rag import (HashingEmbedder, LLMClient, RAGPipeline, VectorStore, chunk_documents, load_bytes,
                 mmr, split_text)
from rag.llm import ChatError, LLMConfig, classify_error
from rag.pipeline import NOT_FOUND, format_context

SAMPLE = os.path.join(os.path.dirname(__file__), "..", "data", "sample")


# ---------- chunking ----------
def test_chunks_respect_size_and_cover_text():
    text = " ".join(f"Sentence number {i} is here." for i in range(200))
    chunks = split_text(text, chunk_size=200, overlap=40)
    assert all(len(c) <= 200 for c in chunks)
    assert len(chunks) > 5
    assert "Sentence number 0" in chunks[0] and "Sentence number 199" in chunks[-1]


def test_overlap_repeats_content_between_neighbours():
    text = " ".join(f"word{i}" for i in range(300))
    chunks = split_text(text, chunk_size=100, overlap=30)
    shared = set(chunks[0].split()) & set(chunks[1].split())
    assert shared, "neighbouring chunks should share words"


def test_short_text_is_one_chunk_and_prefers_paragraphs():
    assert split_text("Just a short note.", 500, 50) == ["Just a short note."]
    chunks = split_text("Para one is here.\n\nPara two is here.", chunk_size=25, overlap=0)
    assert chunks == ["Para one is here.", "Para two is here."]


def test_invalid_chunk_params():
    with pytest.raises(ValueError):
        split_text("x", 100, 100)
    with pytest.raises(ValueError):
        split_text("x", 0, 0)


def test_chunk_ids_unique_and_metadata_kept():
    pages = [("alpha " * 200, {"source": "a.txt", "page": 1}), ("beta " * 200, {"source": "a.txt", "page": 2})]
    chunks = chunk_documents(pages, 200, 20)
    assert len({c.id for c in chunks}) == len(chunks)
    assert chunks[0].metadata["source"] == "a.txt" and chunks[-1].metadata["page"] == 2


# ---------- loaders ----------
def test_load_text_and_unsupported():
    pages = load_bytes("n.txt", b"Hello\r\n\r\n\r\n\r\nWorld")
    assert pages[0][0] == "Hello\n\nWorld" and pages[0][1]["source"] == "n.txt"
    with pytest.raises(ValueError):
        load_bytes("x.exe", b"123")
    with pytest.raises(ValueError):
        load_bytes("empty.txt", b"   \n  ")


def test_load_docx_roundtrip():
    import io
    from docx import Document
    d = Document()
    d.add_paragraph("Quarterly revenue grew twelve percent.")
    d.add_paragraph("Churn is stable.")
    buf = io.BytesIO()
    d.save(buf)
    pages = load_bytes("r.docx", buf.getvalue())
    assert "Quarterly revenue" in pages[0][0] and "Churn" in pages[0][0]


# ---------- embeddings ----------
def test_hashing_embedder_normalised_deterministic_and_related():
    e = HashingEmbedder()
    a, b = e.embed_documents(["annual leave policy for employees", "annual leave days for employees"])
    c = e.embed_query("shipping cost for express delivery")
    assert abs(np.linalg.norm(a) - 1) < 1e-5
    assert (a == e.embed_documents(["annual leave policy for employees"])[0]).all()
    assert a @ b > a @ c


# ---------- vector store ----------
def test_store_add_search_delete_sources():
    e = HashingEmbedder()
    store = VectorStore()
    chunks = chunk_documents([("cats purr and sleep all day", {"source": "cats.txt", "page": 1}),
                              ("rockets burn fuel to reach orbit", {"source": "space.txt", "page": 1})], 100, 0)
    store.add(chunks, e.embed_documents([c.text for c in chunks]))
    assert store.count() == 2 and store.sources() == {"cats.txt": 1, "space.txt": 1}
    top = store.search(e.embed_query("rockets and orbit"), n=2)
    assert top[0]["metadata"]["source"] == "space.txt" and top[0]["score"] > top[1]["score"]
    store.delete_source("cats.txt")
    assert store.sources() == {"space.txt": 1}
    store.reset()
    assert store.count() == 0 and store.search(e.embed_query("x")) == []


def test_collections_are_isolated():
    e = HashingEmbedder()
    s1, s2 = VectorStore(), VectorStore()
    ch = chunk_documents([("hello world", {"source": "a.txt", "page": 1})], 100, 0)
    s1.add(ch, e.embed_documents([c.text for c in ch]))
    assert s1.count() == 1 and s2.count() == 0


# ---------- MMR ----------
def test_mmr_prefers_diverse_over_duplicates():
    q = np.array([1.0, 0.0], dtype="float32")
    cands = np.array([[1.0, 0.0], [0.999, 0.045], [0.6, 0.8]], dtype="float32")
    cands = cands / np.linalg.norm(cands, axis=1, keepdims=True)
    assert mmr(q, cands, k=2, lambda_mult=1.0) == [0, 1]      # pure relevance picks the near-duplicate
    assert mmr(q, cands, k=2, lambda_mult=0.3) == [0, 2]      # diversity picks the different passage
    assert mmr(q, np.zeros((0, 2), "float32"), k=2) == []


# ---------- pipeline ----------
def build_pipeline(llm=None, size=500):
    pipe = RAGPipeline(VectorStore(), HashingEmbedder(), llm, chunk_size=size, overlap=80)
    from rag import load_path
    for p in sorted(glob.glob(os.path.join(SAMPLE, "*"))):
        pipe.ingest_pages(load_path(p))
    return pipe


def test_retrieval_finds_answer_bearing_chunk():
    pipe = build_pipeline()
    hits = pipe.retriever.retrieve("How many days of paid annual leave do I get?", k=3)
    assert any("24 days of paid annual leave" in h.text for h in hits)
    assert hits[0].source == "employee_handbook.md"


def test_threshold_filters_and_no_context_answer():
    pipe = build_pipeline()
    assert pipe.ask("zebra quantum xylophone", threshold=0.9).mode == "no-context"
    assert pipe.ask("zebra quantum xylophone", threshold=0.9).text == NOT_FOUND


def test_reindexing_same_file_replaces_chunks():
    pipe = RAGPipeline(VectorStore(), HashingEmbedder())
    pipe.ingest_file("a.txt", b"old content about apples " * 30)
    first = pipe.store.count()
    pipe.ingest_file("a.txt", b"old content about apples " * 30)
    assert pipe.store.count() == first


def test_retrieval_only_mode_without_llm():
    ans = build_pipeline(llm=None).ask("What is the minimum password length?")
    assert ans.mode == "retrieval-only" and ans.hits


def test_generation_prompt_is_grounded_and_uses_history():
    seen = {}

    def fake(messages, temperature, max_tokens):
        seen["messages"] = messages
        return "You get 24 days [1]."

    pipe = build_pipeline(llm=LLMClient(transport=fake))
    history = [{"role": "user", "content": "hi"}, {"role": "assistant", "content": "hello"}]
    ans = pipe.ask("How many days of annual leave?", history=history, k=3)
    msgs = seen["messages"]
    assert ans.mode == "generated" and ans.text.startswith("You get 24")
    assert msgs[0]["role"] == "system" and "ONLY" in msgs[0]["content"] and NOT_FOUND in msgs[0]["content"]
    assert msgs[1]["content"] == "hi"                       # history included
    assert "[1] (source:" in msgs[-1]["content"] and "Question:" in msgs[-1]["content"]


def test_format_context_numbers_sources():
    from rag import Hit
    ctx = format_context([Hit("1", "text a", "a.pdf", 3, 0.9), Hit("2", "text b", "b.md", 1, 0.8)])
    assert ctx.startswith("[1] (source: a.pdf, page 3)") and "[2] (source: b.md, page 1)" in ctx


# ---------- LLM client ----------
class HTTPErr(Exception):
    def __init__(self, status):
        super().__init__(f"status {status}")
        self.status_code = status


def test_llm_retries_then_succeeds_and_fails_fast_on_auth():
    outcomes = iter([HTTPErr(429), HTTPErr(503), "ok"])
    sleeps = []

    def t(m, temp, mx):
        o = next(outcomes)
        if isinstance(o, Exception):
            raise o
        return o

    assert LLMClient(transport=t, sleep=sleeps.append).chat([]) == "ok" and len(sleeps) == 2

    def bad(m, temp, mx):
        raise HTTPErr(401)

    with pytest.raises(ChatError) as e:
        LLMClient(transport=bad, sleep=sleeps.append).chat([])
    assert "API key" in e.value.user_message and not e.value.retryable


def test_config_none_without_key(monkeypatch):
    monkeypatch.delenv("LLM_API_KEY", raising=False)
    assert LLMConfig.from_env({}) is None
    assert LLMConfig.from_env({"LLM_API_KEY": " k "}).api_key == "k"
    assert "Provider said" in classify_error(HTTPErr(404)).user_message
